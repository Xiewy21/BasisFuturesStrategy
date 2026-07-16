"""report.py —— （可选）用 python-docx 生成汇总报告。

方法论性质：无（纯展示层，不产出任何新的模型结论），只是把 run.py 跑出来的
敏感度表 / 回测指标排版成一份 Word 文档，方便分享。python-docx 是可选依赖，
未安装时本模块的函数会打印提示并跳过，不影响四层流水线本身的运行。
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

import pandas as pd


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401

        return True
    except ImportError:
        return False


def generate_report(
    output_path: str | Path,
    sensitivity_df: pd.DataFrame,
    backtest_metrics: dict,
    factor_summary: Optional[pd.DataFrame] = None,
    title: str = "股指期货基差研究报告（自动生成，仅供研究参考）",
) -> Optional[Path]:
    """生成一份简单的 Word 汇总报告。python-docx 未安装时跳过并返回 None。

    Args:
        output_path: 输出 .docx 路径。
        sensitivity_df: forecast.sensitivity_table 的输出。
        backtest_metrics: strategy.backtest_timing 的输出（不含 equity_curve 也可）。
        factor_summary: 可选，因子面板的 describe() 之类的统计摘要。
        title: 报告标题。

    Returns:
        实际写出的文件路径；若 python-docx 未安装则返回 None。
    """
    if not _docx_available():
        print("[report] 未安装 python-docx，跳过报告生成（pip install python-docx 可启用）。")
        return None

    import docx

    doc = docx.Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(
        "本报告由 basis_model.report 自动生成，属于研究/回测工具输出，"
        "不构成投资建议，不能直接用于实盘交易决策。"
    )

    doc.add_heading("一、情景敏感度表（预测层输出）", level=2)
    if sensitivity_df is not None and not sensitivity_df.empty:
        table = doc.add_table(rows=1, cols=len(sensitivity_df.columns))
        table.style = "Light Grid Accent 1"
        hdr_cells = table.rows[0].cells
        for i, col in enumerate(sensitivity_df.columns):
            hdr_cells[i].text = str(col)
        for _, row in sensitivity_df.iterrows():
            cells = table.add_row().cells
            for i, col in enumerate(sensitivity_df.columns):
                cells[i].text = f"{row[col]}"
    else:
        doc.add_paragraph("（无数据）")

    doc.add_heading("二、择时规则回测结果（规则回测，非模型）", level=2)
    for k, v in backtest_metrics.items():
        if k == "equity_curve":
            continue
        doc.add_paragraph(f"{k}: {v}")

    if factor_summary is not None and not factor_summary.empty:
        doc.add_heading("三、因子面板统计摘要", level=2)
        table = doc.add_table(rows=1, cols=len(factor_summary.columns) + 1)
        table.style = "Light Grid Accent 1"
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "stat"
        for i, col in enumerate(factor_summary.columns):
            hdr_cells[i + 1].text = str(col)
        for idx, row in factor_summary.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(idx)
            for i, col in enumerate(factor_summary.columns):
                cells[i + 1].text = f"{row[col]:.4f}" if pd.notna(row[col]) else "NaN"

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"[report] 报告已保存: {out_path}")
    return out_path
